"""
One consistent look for every Excel export (openpyxl):

    Anchorstone Capital Ltd — <Report title>
    <filter / as-of lines>
    (blank)
    [navy header row, frozen, with filters]
    rows...
    [bold totals row]

Money cells are real numbers formatted #,##0.00 so the spreadsheet can add them
up; dates are real dates. Timestamps are shown in Nairobi time.
"""
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils import brand
from app.utils.dates import EAT

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_NAVY = brand.PRIMARY.lstrip("#").upper()
_TINT = brand.SURFACE_ALT.lstrip("#").upper()
FORMATS = {
    "money": "#,##0.00",
    "int": "#,##0",
    "date": "dd mmm yyyy",
    "datetime": "dd mmm yyyy hh:mm",
    "percent": "0.00%",
    "rate": '0.00"%"',
    "text": "@",
}


@dataclass
class Col:
    header: str
    width: int = 16
    kind: str = "text"  # text | money | int | date | datetime | percent | rate


def _cell_value(value, kind: str):
    if value is None:
        return None
    if kind in ("money", "rate") and not isinstance(value, (int, float)):
        return float(Decimal(value))
    if kind == "percent" and value is not None:
        return float(value)
    if kind == "datetime" and isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(EAT).replace(tzinfo=None)
        return value
    if kind == "date" and isinstance(value, datetime):
        return value.date()
    return value


def build_workbook(title: str, info_lines: list[str], cols: list[Col], rows: list[list],
                   totals: dict | None = None, sheet_name: str = "Report") -> bytes:
    """totals = {"label": "Total", <column index>: value, ...}; label goes in the first column."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    ws.cell(row=1, column=1, value=f"{brand.COMPANY_NAME} — {title}").font = Font(bold=True, size=14, color=_NAVY)
    for i, line in enumerate(info_lines, start=2):
        ws.cell(row=i, column=1, value=line).font = Font(size=10, color="6C7088")
    header_row = len(info_lines) + 3

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=_NAVY)
    for c, col in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=c, value=col.header)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="right" if col.kind in ("money", "int", "percent", "rate") else "left",
                                   vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = col.width
    ws.row_dimensions[header_row].height = 20

    r = header_row
    for r, row in enumerate(rows, start=header_row + 1):
        for c, (value, col) in enumerate(zip(row, cols), start=1):
            cell = ws.cell(row=r, column=c, value=_cell_value(value, col.kind))
            if col.kind != "text" and value is not None:
                cell.number_format = FORMATS[col.kind]
    last_data_row = r

    if not rows:
        ws.cell(row=header_row + 1, column=1, value="No records match these filters.").font = Font(italic=True, color="6C7088")
        last_data_row = header_row + 1

    if totals:
        tr = last_data_row + 1
        top = Border(top=Side(style="thin", color=_NAVY))
        fill = PatternFill("solid", fgColor=_TINT)
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=tr, column=c)
            cell.border, cell.fill, cell.font = top, fill, Font(bold=True, color=_NAVY)
            value = totals.get("label") if c == 1 else totals.get(c - 1)
            if value is not None:
                cell.value = _cell_value(value, "text" if c == 1 else col.kind)
                if c > 1 and col.kind != "text":
                    cell.number_format = FORMATS[col.kind]

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{last_data_row}"
    ws.sheet_view.showGridLines = False if not rows else True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fmt_range(start: date | None, end: date | None) -> str:
    if start and end:
        return f"{start:%d %b %Y} to {end:%d %b %Y}"
    if start:
        return f"From {start:%d %b %Y}"
    if end:
        return f"Up to {end:%d %b %Y}"
    return "All dates"
