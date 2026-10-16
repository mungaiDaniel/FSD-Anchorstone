"""
Golden files for the investor → party migration (phase 2b item 2).

The golden files were made ONCE, by tests/golden/test_make_golden.py run at the
last commit before phase 2b (8cebb7e), so they record what the old code showed:

    GOLDEN_WRITE=1 venv/Scripts/python -m pytest tests/golden/test_make_golden.py

    pre_2b_data.json.gz     the seeded rows (demo book, May–Sep 2026)
    pre_2b_outputs.json.gz  everything staff and investors see from them

tests/test_investor_migration.py loads the rows into the pre-2b schema, runs
the migrations (which create parties and accounts), captures the same outputs
with the current code and requires them to be identical. Only wall-clock
"Generated …" stamps are left out.
"""
import base64
import gzip
import io
import json
import re
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

from sqlalchemy import text

HERE = Path(__file__).parent
DATA_FILE = HERE / "pre_2b_data.json.gz"
OUTPUTS_FILE = HERE / "pre_2b_outputs.json.gz"
PRE_2B_REVISION = "a6f7b8c9d0e1"
TODAY = date(2026, 10, 5)

STAFF_EMAIL = "staff@anchorstone.test"          # records deposits and requests
CHECKER_A = "admin1@anchorstone.test"           # approves withdrawals, calculates interest
CHECKER_B = "admin2@anchorstone.test"           # approves interest
SUPER = "super@anchorstone.test"
EXCLUDED_TABLES = {"alembic_version"}

REPORT_PARAMS = {
    "audit-log": {"from": "2026-01-01", "to": "2026-12-31"},  # first, before exports audit themselves
    "investors": {},
    "deposits": {"from": "2026-05-01", "to": "2026-10-05"},
    "withdrawals": {"from": "2026-05-01", "to": "2026-10-05"},
    "interest": {},
    "aum": {},
    "email-log": {"from": "2026-01-01", "to": "2026-12-31"},
}
_GENERATED_PDF = re.compile(r"Generated \d{1,2} \w+ \d{4} \d{2}:\d{2} UTC")


# ── Data fixture ───────────────────────────────────────────────────────────


def dump_data(db, skip_tables) -> dict:
    """Every non-reference table with rows, as COPY CSV, plus sequence positions."""
    raw = db.session.connection().connection.dbapi_connection
    tables = [r[0] for r in db.session.execute(text(
        "SELECT tablename FROM pg_tables WHERE schemaname = 'public' ORDER BY tablename"))]
    out = {"tables": {}, "sequences": {}}
    for name in tables:
        if name in skip_tables or name in EXCLUDED_TABLES:
            continue
        if not db.session.execute(text(f'SELECT EXISTS (SELECT 1 FROM "{name}")')).scalar():
            continue
        buf = io.StringIO()
        with raw.cursor() as cur:
            cur.copy_expert(f'COPY (SELECT * FROM "{name}" ORDER BY 1) TO STDOUT WITH CSV HEADER', buf)
        out["tables"][name] = buf.getvalue()
    for (seq,) in db.session.execute(text("SELECT sequencename FROM pg_sequences WHERE schemaname = 'public'")):
        value = db.session.execute(text(f'SELECT last_value, is_called FROM "{seq}"')).one()
        if value.is_called:
            out["sequences"][seq] = value.last_value
    return out


def load_data(db, data: dict, table_order) -> None:
    """Insert the dumped rows (parents first) and put the sequences back."""
    raw = db.session.connection().connection.dbapi_connection
    names = [n for n in table_order if n in data["tables"]]
    missing = set(data["tables"]) - set(names)
    assert not missing, f"tables in the fixture but not in the schema: {missing}"
    for name in names:
        csv_text = data["tables"][name]
        header = csv_text.split("\n", 1)[0]
        with raw.cursor() as cur:
            cur.copy_expert(f'COPY "{name}" ({header}) FROM STDIN WITH CSV HEADER', io.StringIO(csv_text))
    for seq, value in data["sequences"].items():
        db.session.execute(text("SELECT setval(:s, :v, true)"), {"s": seq, "v": value})


def write(path: Path, payload) -> None:
    path.write_bytes(gzip.compress(json.dumps(payload, sort_keys=True, indent=1).encode(), mtime=0))


def read(path: Path):
    return json.loads(gzip.decompress(path.read_bytes()))


# ── Outputs ────────────────────────────────────────────────────────────────


def pdf_text(pdf: bytes) -> str:
    """Text drawn on the pages (ReportLab: ASCII85 + Flate streams of (...) Tj)."""
    import base64 as b64
    import zlib

    out = []
    for raw in re.findall(rb"stream\r?\n(.*?)~>\s*endstream", pdf, re.S):
        try:
            raw = zlib.decompress(b64.a85decode(raw, adobe=False, ignorechars=b" \t\r\n"))
        except (ValueError, zlib.error):
            continue
        out += [m.decode("latin-1") for m in re.findall(rb"\((.*?)\)\s*Tj", raw)]
    return _GENERATED_PDF.sub("Generated <now>", " ".join(out))


def _cells(xlsx: bytes) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx))
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else (v.isoformat() if isinstance(v, (date, datetime)) else repr(v))
                      for v in row]
            if values and values[0].startswith("'Generated "):
                continue  # wall-clock stamp
            rows.append(values)
        sheets[ws.title] = rows
    return sheets


def capture(client, headers_for) -> dict:
    """Everything the old screens, statements and exports show. headers_for(email) signs a request."""
    from app.utils.pdf_generator import generate_account_statement_pdf, generate_interest_statement_pdf

    staff, admin = headers_for(STAFF_EMAIL), headers_for(SUPER)

    def get(url, headers=staff):
        res = client.get(url, headers=headers)
        assert res.status_code == 200, (url, res.status_code, res.get_data(as_text=True)[:300])
        return res

    out = {}
    for name, params in REPORT_PARAMS.items():
        query = "&".join(f"{k}={v}" for k, v in params.items())
        out[f"report:{name}"] = _cells(get(f"/api/v1/reports/{name}.xlsx?{query}", admin).data)

    for status in ("", "active", "closed"):
        out[f"GET /investors?status={status}"] = get(f"/api/v1/investors?per_page=200&status={status}").get_json()
    out["GET /investors?search=wan"] = get("/api/v1/investors?per_page=200&search=wan").get_json()
    ids = [i["id"] for i in out["GET /investors?status="]["data"]["items"]]
    for i in sorted(ids):
        for suffix in ("", "/ledger", "/deposits", "/withdrawals", "/payouts", "/emails"):
            out[f"GET /investors/{i}{suffix}"] = get(f"/api/v1/investors/{i}{suffix}").get_json()
        out[f"pdf:account-statement/{i}"] = pdf_text(generate_account_statement_pdf(i))
        for p in out[f"GET /investors/{i}/payouts"]["data"]:
            out[f"pdf:interest-statement/{p['id']}"] = pdf_text(generate_interest_statement_pdf(p["id"]))
            out[f"GET statement.pdf {p['id']}"] = base64.b64encode(
                get(f"/api/v1/interest/payouts/{p['id']}/statement.pdf").data).decode()
    for url in ("/api/v1/deposits?per_page=200", "/api/v1/withdrawals?per_page=200", "/api/v1/emails?per_page=200",
                "/api/v1/interest/periods", "/api/v1/stats/overview"):
        out[f"GET {url}"] = get(url).get_json()
    return out


# ── The demo book (a copy of scripts/seed_demo.py with neutral addresses) ─


INVESTORS = [  # key, full name, phone, ID, KRA PIN, payout method, bank, account
    ("grace", "Grace Wanjiku Mwangi", "0712 345 678", "28456712", "A012345678K", "mpesa", None, "0712345678"),
    ("peter", "Peter Kamau Njoroge", "0722 118 904", "22981045", "A009876543P", "bank", "KCB Bank", "1145829364"),
    ("aisha", "Aisha Hassan Abdi", "0733 456 120", "31209876", "A015678234M", "mpesa", None, "0733456120"),
    ("john", "John Ochieng Otieno", "0701 223 344", "25678901", "A011223344J", "bank", "Equity Bank", "0180293847561"),
    ("mary", "Mary Achieng Owino", "0745 667 889", "33456789", "A018765432Q", "mpesa", None, "0745667889"),
    ("david", "David Kiprotich Rotich", "0710 998 877", "20345678", "A007654321D", "bank", "Co-operative Bank", "01109283746500"),
    ("fatuma", "Fatuma Said Ali", "0798 112 233", "29871234", "A013579246F", "bank", "NCBA Bank", "1002938475"),
    ("james", "James Mutua Musyoka", "0723 554 466", "24567123", "A010293847R", "bank", "NCBA Bank", "1004857362"),
    ("esther", "Esther Nyambura Kimani", "0757 889 900", "34567890", "A019283746E", "mpesa", None, "0757889900"),
    ("samuel", "Samuel Kipchoge Langat", "0711 445 566", "21234567", "A008172635S", "bank", "Stanbic Bank", "0100003829471"),
    ("lucy", "Lucy Wairimu Gitau", "0768 223 445", "35678123", "A020394857L", "mpesa", None, "0768223445"),
    # Edge cases the migration must carry through unchanged:
    ("hassan", "Hassan Omar Juma", "+254 790 334 556", "26789012", "a012938475h", "bank", "Equity Bank", "0180574839201"),
    ("kamau", "Kamau", None, None, None, None, None, None),                      # one name, nothing else
    ("closed", "Wambui Closed-Account", "not a phone", "A1234567", None, "mpesa", None, "0700111222"),  # passport-like
]


def seed_demo_book(db) -> None:
    """Run the demo book through the (pre-2b) services, with dates pinned like scripts/seed_demo.py."""
    from scripts.seed_demo import EVENTS
    from app.Emails.model import STATUS_QUEUED, STATUS_SENT, OutboxEmail
    from app.Interest import service as interest
    from app.Interest.model import InterestPayout, InterestPeriod, InterestRate
    from app.Investors.service import create_investor, set_investor_status
    from app.Transactions import service as tx
    from app.Transactions.model import Deposit, Withdrawal
    from app.utils.dates import month_bounds

    eat = timezone(timedelta(hours=3))

    def at(day, hour=10, minute=0):
        return datetime.combine(day, time(hour, minute), eat).astimezone(timezone.utc)

    d = date.fromisoformat
    interest.set_rate({"rate_percent": "11", "effective_from": "2026-05-01", "note": "Launch rate"}, CHECKER_A)
    interest.set_rate({"rate_percent": "12", "effective_from": "2026-08-01", "note": "Board-approved increase"}, CHECKER_A)
    for r in InterestRate.query.all():
        r.created_at = at(r.effective_from - timedelta(days=3), 15)

    first_dep = {}
    for e in EVENTS:
        if e[0] == "dep":
            first_dep.setdefault(e[2], d(e[1]))
    inv = {}
    for key, name, phone, id_no, pin, method, bank, account in INVESTORS:
        i = create_investor({"full_name": name, "email": f"investor+{key}@example.com", "phone": phone,
                             "id_number": id_no, "kra_pin": pin, "payout_method": method, "bank_name": bank,
                             "payout_account": account}, STAFF_EMAIL)
        i.created_at = at(first_dep.get(key, date(2026, 9, 20)), 9, 30)
        inv[key] = i.id
    db.session.commit()
    set_investor_status(inv["closed"], {"status": "closed"}, CHECKER_A)

    for e in EVENTS:
        kind = e[0]
        if kind == "dep":
            _, day, key, amount, method, ref = e
            dep = tx.record_deposit({"investor_id": inv[key], "amount": amount, "value_date": day,
                                     "method": method, "reference": ref}, STAFF_EMAIL)
            dep.created_at = at(d(day), 11)
        elif kind == "rev":
            dep = Deposit.query.filter_by(reference=e[1]).one()
            tx.reverse_deposit(dep.id, {"reason": e[2]}, CHECKER_A)
            dep.reversed_at = at(dep.value_date, 16)
        elif kind == "wd":
            _, day, key, amount, method, decision = e
            w = tx.request_withdrawal({"investor_id": inv[key], "amount": amount, "requested_date": day,
                                       "method": method}, STAFF_EMAIL)
            w.created_at = at(d(day), 12)
            if decision == "pending":
                pass
            elif isinstance(decision, tuple):
                tx.reject_withdrawal(w.id, {"reason": decision[1]}, CHECKER_A)
                w.decided_at = at(d(day) + timedelta(days=1), 10)
            else:
                ref = f"EFT{d(decision):%y%m%d}{w.id:03d}" if method == "bank_transfer" else f"SQ{w.id:02d}{d(decision):%m%d}ZK"
                tx.approve_withdrawal(w.id, {"value_date": decision, "payment_reference": ref}, CHECKER_A)
                w.decided_at = at(d(decision), 14)
        elif kind == "month":
            _, year, month = e
            period = interest.calculate_period({"year": year, "month": month}, CHECKER_A)
            interest.approve_period(period.id, CHECKER_B)
            payouts = period.payouts.all()
            interest.mark_payouts_paid(period.id, {"payments": [
                {"payout_id": p.id, "reference": (f"SQ{p.id:03d}{month:02d}PAY" if p.investor.payout_method == "mpesa"
                                                  else f"EFT{year % 100}{month:02d}{p.id:04d}")}
                for p in payouts]}, CHECKER_A)
            _, end = month_bounds(year, month)
            period.calculated_at = at(end + timedelta(days=2), 9)
            period.approved_at = at(end + timedelta(days=2), 15)
            period.paid_at = at(end + timedelta(days=5), 11)
            for p in payouts:
                p.paid_at = period.paid_at
        db.session.commit()

    for email in OutboxEmail.query.filter(OutboxEmail.status == STATUS_QUEUED).all():
        when = None
        if email.deposit_id:
            when = db.session.get(Deposit, email.deposit_id).created_at + timedelta(minutes=3)
        elif email.withdrawal_id:
            when = db.session.get(Withdrawal, email.withdrawal_id).decided_at + timedelta(minutes=3)
        elif email.payout_id:
            when = db.session.get(InterestPayout, email.payout_id).period.approved_at + timedelta(minutes=5)
        if when is not None:
            email.queued_at = when
            email.status, email.sent_at, email.attempts = STATUS_SENT, when + timedelta(minutes=2), 1
    db.session.commit()
    assert InterestPeriod.query.count() == 4
