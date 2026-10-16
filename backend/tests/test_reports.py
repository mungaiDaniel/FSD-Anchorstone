import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.Emails.model import OutboxEmail
from app.Emails.sender import drain_queue, mail
from app.Reports.scheduled import scheduled_run
from app.utils.audit_log import AuditLog


def sheet(res):
    assert res.status_code == 200, res.get_data(as_text=True)[:300]
    assert res.headers["Content-Type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(res.data)).active


def table(ws):
    """Rows from the navy header row down, header first."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    start = next(i for i, r in enumerate(rows) if i > 1 and r[0] and len(r) > 1 and r[1])
    return rows[start:]


@pytest.fixture
def book(make_investor, deposit, withdraw, rate, admin1, admin2):
    rate("12", "2026-01-01")
    a = make_investor(full_name="Aisha Hassan")
    b = make_investor(full_name="Brian Kiprop")
    deposit(a["id"], "100000", "2026-08-03")
    deposit(b["id"], "50000", "2026-08-20")
    deposit(a["id"], "20000", "2026-09-10")
    withdraw(b["id"], "10000", "2026-09-15")
    pid = admin1.post("/api/v1/interest/periods/calculate", {"year": 2026, "month": 8}).get_json()["data"]["id"]
    admin2.post(f"/api/v1/interest/periods/{pid}/approve")
    return {"a": a, "b": b, "period": pid}


# ── Permissions ──


def test_staff_can_export_financial_reports_not_logs(staff, book):
    for name in ("investors", "deposits", "withdrawals", "interest", "aum"):
        assert staff.get(f"/api/v1/reports/{name}.xlsx").status_code == 200, name
    assert staff.get("/api/v1/reports/audit-log.xlsx").status_code == 403
    assert staff.get("/api/v1/reports/email-log.xlsx").status_code == 403
    names = [r["name"] for r in staff.get("/api/v1/reports").get_json()["data"]]
    assert "audit-log" not in names and "deposits" in names


def test_admins_can_export_logs_and_view_audit_screen(admin1, book):
    assert admin1.get("/api/v1/reports/audit-log.xlsx").status_code == 200
    assert admin1.get("/api/v1/reports/email-log.xlsx").status_code == 200
    assert admin1.get("/api/v1/audit-logs").status_code == 200


def test_unknown_report_and_no_login(client, staff):
    assert staff.get("/api/v1/reports/nope.xlsx").status_code == 404
    assert client.get("/api/v1/reports/deposits.xlsx").status_code == 401


# ── Content ──


def test_investors_report(staff, book):
    rows = table(sheet(staff.get("/api/v1/reports/investors.xlsx")))
    header, body, total = rows[0], rows[1:-1], rows[-1]
    assert header[:2] == ["Client code", "Name"]
    by_name = {r[1]: r for r in body}
    assert by_name["Aisha Hassan"][5] == 120000 and by_name["Brian Kiprop"][5] == 40000
    assert by_name["Brian Kiprop"][7] == 10000
    assert total[5] == 160000


def test_deposits_report_date_range_and_totals(staff, book):
    rows = table(sheet(staff.get("/api/v1/reports/deposits.xlsx?from=2026-08-01&to=2026-08-31")))
    body, total = rows[1:-1], rows[-1]
    assert len(body) == 2
    assert total[5] == 150000 and "2 deposits" in total[0]


def test_withdrawals_report_totals_approved_only(staff, book):
    staff.post("/api/v1/withdrawals", {"investor_id": book["a"]["id"], "amount": "5000", "method": "mpesa"})
    rows = table(sheet(staff.get("/api/v1/reports/withdrawals.xlsx")))
    body, total = rows[1:-1], rows[-1]
    assert sorted(r[7] for r in body) == ["Approved", "Pending"]
    assert total[6] == 10000


def test_interest_report_per_month(staff, book):
    rows = table(sheet(staff.get(f"/api/v1/reports/interest.xlsx?period_id={book['period']}")))
    body, total = rows[1:-1], rows[-1]
    assert {r[0] for r in body} == {"August 2026"}
    assert len(body) == 2
    assert total[8] == pytest.approx(sum(r[8] for r in body))
    assert body[0][13] == "admin2@anchorstone.test"  # approved by


def test_aum_report_month_over_month(staff, book):
    rows = table(sheet(staff.get("/api/v1/reports/aum.xlsx")))[1:]
    by_month = {r[0].split(" (")[0]: r for r in rows}
    assert by_month["August 2026"][2] == 150000
    assert by_month["September 2026"][2] == 160000
    assert by_month["September 2026"][3] == 10000
    assert by_month["September 2026"][4] == pytest.approx(10000 / 150000)
    assert "October 2026" in rows[-1][0]  # current month, to date


def test_every_export_is_audited_and_audit_export_filters_by_user(admin1, staff, book):
    staff.get("/api/v1/reports/deposits.xlsx?from=2026-08-01&to=2026-08-31")
    log = AuditLog.query.filter_by(action="EXPORT_REPORT", target_name="deposits").one()
    assert log.user_email == "staff@anchorstone.test"
    assert "01 Aug 2026 to 31 Aug 2026" in log.description
    rows = table(sheet(admin1.get("/api/v1/reports/audit-log.xlsx?user=staff@")))[1:]
    assert rows and all(r[2] == "staff@anchorstone.test" for r in rows)


def test_email_log_export(admin1, book):
    rows = table(sheet(admin1.get("/api/v1/reports/email-log.xlsx")))[1:]
    assert any(r[4] == "Deposit confirmation" for r in rows)


def test_bad_date_range(staff):
    assert staff.get("/api/v1/reports/deposits.xlsx?from=2026-09-30&to=2026-09-01").status_code == 400


# ── Scheduled AUM summary ──


def test_schedule_settings_super_admin_only(admin1, superadmin):
    assert admin1.get("/api/v1/settings/scheduled-reports").status_code == 403
    assert superadmin.patch("/api/v1/settings/scheduled-reports", {"aum_enabled": True}).status_code == 400
    res = superadmin.patch("/api/v1/settings/scheduled-reports",
                           {"aum_enabled": True, "aum_recipients": "md@anchorstone.test, Finance@Anchorstone.test"})
    assert res.status_code == 200
    assert res.get_json()["data"]["aum_recipients"] == ["md@anchorstone.test", "finance@anchorstone.test"]
    assert superadmin.patch("/api/v1/settings/scheduled-reports", {"aum_recipients": "not-an-email"}).status_code == 400


def test_scheduled_run_creates_pending_emails_once(superadmin, admin1, book):
    superadmin.patch("/api/v1/settings/scheduled-reports",
                     {"aum_enabled": True, "aum_recipients": "md@anchorstone.test, cfo@anchorstone.test"})
    assert scheduled_run(date(2026, 10, 1)) == 2
    assert scheduled_run(date(2026, 10, 1)) == 0  # the same month is never generated twice
    emails = OutboxEmail.query.filter_by(email_type="AUM_REPORT").all()
    assert {e.status for e in emails} == {"pending"} and all(e.investor_id is None for e in emails)
    assert emails[0].subject == "Monthly AUM summary — September 2026"
    assert "KES 160,000.00" in emails[0].body
    assert emails[0].attachment.filename.endswith(".xlsx")

    # Nothing goes out until an admin confirms; then each has the spreadsheet attached.
    with mail.record_messages() as outbox:
        drain_queue()
    assert not [m for m in outbox if "AUM" in m.subject]
    admin1.post("/api/v1/emails/confirm", {"ids": [e.id for e in emails]})
    with mail.record_messages() as outbox:
        drain_queue()
    aum = [m for m in outbox if "AUM" in m.subject]
    assert len(aum) == 2 and aum[0].attachments[0].filename.endswith(".xlsx")
    assert admin1.get("/api/v1/emails?type=AUM_REPORT").status_code == 200


def test_scheduled_run_does_nothing_when_off(superadmin, book):
    superadmin.patch("/api/v1/settings/scheduled-reports", {"aum_enabled": False, "aum_recipients": "md@anchorstone.test"})
    assert scheduled_run(date(2026, 10, 1)) == 0


def test_run_now(superadmin, book):
    superadmin.patch("/api/v1/settings/scheduled-reports", {"aum_recipients": "md@anchorstone.test"})
    res = superadmin.post("/api/v1/settings/scheduled-reports/run-now")
    assert res.status_code == 201
    assert OutboxEmail.query.filter_by(email_type="AUM_REPORT", status="pending").count() == 1
